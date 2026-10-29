"""Assemble the pitch deck from the slide images rendered by Marketing/video (Deck composition).

    python build_pptx.py <slides dir> <out.pptx>

Each slide is one full-bleed image with a fade transition and the speaker notes below.
Keynote opens the .pptx directly; the wording comes from docs/pitch/script.md.
"""
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Emu

NOTES = [
    "大家好，我们是 PRTS。",
    "在中国，有 1731 万视障人士。",
    "在岗的导盲犬，只有 410 多只。",
    "大多数人手里只有一根导盲杖。它能告诉你脚下一米，却说不出前面是什么、该往哪走。",
    "现有的电子方案，要么贵，要么离不开网络。",
    "我们想要的，是一部 iPhone 就够。",
    "PRTS 做三件事：避开障碍、描述画面、读出文字。",
    "打开避障，激光雷达每秒测十次，找出地面和障碍，规划一条能走的路，再用语音和左右震动带你走。",
    "前方空旷时只做播报；障碍远时先引导走近；很近时逐点绕过去。",
    "分析每秒十次，95% 的结果在 48 毫秒内送到，整个避障不联网。",
    "短按屏幕，描述眼前的画面；三根手指按住，读出纸上的字。读字也在手机上完成。",
    "我们在夜晚的街道上实地走过、录下 2128 帧数据反复重放。对视障用户，报错和漏报一样危险，我们两样都盯着。",
    "接下来，我们要把视觉模型放进手机，加上台阶检测，并请视障朋友一起实测。",
    "PRTS，让每一步都有回应。谢谢。",
]

FADE = '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="med"><p:fade/></p:transition>'

slides_dir, out = Path(sys.argv[1]), sys.argv[2]
deck = Presentation()
deck.slide_width, deck.slide_height = Emu(12192000), Emu(6858000)
blank = deck.slide_layouts[6]
for i, note in enumerate(NOTES, 1):
    slide = deck.slides.add_slide(blank)
    slide.shapes.add_picture(str(slides_dir / f"slide-{i:02d}.png"), 0, 0, deck.slide_width, deck.slide_height)
    slide.notes_slide.notes_text_frame.text = note
    # The transition element goes right after <p:clrMapOvr> (or <p:cSld>) in the slide XML.
    root = slide._element
    anchor = root.find('{http://schemas.openxmlformats.org/presentationml/2006/main}clrMapOvr')
    (anchor if anchor is not None else root[0]).addnext(etree.fromstring(FADE))
deck.save(out)
print(out, len(NOTES), "slides")
